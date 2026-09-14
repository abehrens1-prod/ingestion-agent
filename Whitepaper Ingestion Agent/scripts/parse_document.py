"""
parse_document.py — Convert PDF or DOCX whitepaper to structured tagged plaintext.

Outputs the same tagged format as the blog agent's DOCX walker:
  [H1], [H2], [H3], [P], [BOLD], [LIST], [TABLE_START NxM], [TABLE_END], [BOX_START], [BOX_END]

PDF mode uses pdfplumber with font-size-based heading detection.
DOCX mode adapts the blog agent's paragraph walker.

Usage:
    python scripts/parse_document.py input/whitepaper.pdf
    python scripts/parse_document.py input/whitepaper.docx
    python scripts/parse_document.py input/whitepaper.pdf --output output/source.txt
"""

import re
import sys
import logging
import argparse
import statistics
from pathlib import Path
from typing import Optional

PROJECT_ROOT = Path(__file__).resolve().parent.parent
REPOSITORY_ROOT = PROJECT_ROOT.parent
if str(REPOSITORY_ROOT) not in sys.path:
    sys.path.insert(0, str(REPOSITORY_ROOT))

from ingestion_common.console import enable_utf8

enable_utf8()

logger = logging.getLogger(__name__)


# ---------------------------------------------------------------------------
# PDF OCR fallback (Claude Vision for image-based PDFs)
# ---------------------------------------------------------------------------

_OCR_PROMPT = """\
You are extracting text from whitepaper PDF pages for a content pipeline.

Output ONLY tagged lines — no page separators, no commentary, no markdown.

Tag rules:
- [H1] main document title (appears once, on the first page it occurs)
- [H2] major section headings
- [H3] subsection headings
- [BOLD] bold/emphasized lines that are not headings
- [P] regular body paragraphs
- [LIST] bullet points and numbered list items — one tag per item, strip the bullet/number prefix
- [TABLE_START NxM] followed by "  | cell | cell |" rows, then [TABLE_END]
- [BOX_START] / content lines / [BOX_END] for callout boxes or sidebars

Skip: page numbers, running headers/footers, decorative rules, image captions with no content value.
Preserve ALL numbered heading prefixes exactly (e.g. "1. Section Title" not "Section Title").
"""


def _extract_pdf_ocr(pdf_path: str) -> str:
    """Render PDF pages as images and extract tagged text via Claude Vision."""
    try:
        import fitz  # pymupdf
    except ImportError:
        raise ImportError("pymupdf not installed. Run: pip3 install pymupdf")

    import base64
    import anthropic
    from dotenv import load_dotenv
    load_dotenv(Path(__file__).parent.parent / ".env")

    client = anthropic.Anthropic()

    doc = fitz.open(pdf_path)
    num_pages = len(doc)
    matrix = fitz.Matrix(150 / 72, 150 / 72)  # 150 DPI

    page_b64 = []
    for i in range(num_pages):
        pix = doc[i].get_pixmap(matrix=matrix, colorspace=fitz.csRGB)
        page_b64.append(base64.standard_b64encode(pix.tobytes("png")).decode("ascii"))
    doc.close()

    BATCH = 8
    all_lines = []

    for batch_start in range(0, num_pages, BATCH):
        batch_end = min(batch_start + BATCH, num_pages)
        logger.info(f"OCR: pages {batch_start + 1}–{batch_end} of {num_pages}...")

        content = []
        for idx, b64 in enumerate(page_b64[batch_start:batch_end]):
            content.append({"type": "text", "text": f"Page {batch_start + idx + 1}:"})
            content.append({
                "type": "image",
                "source": {"type": "base64", "media_type": "image/png", "data": b64},
            })
        content.append({"type": "text", "text": _OCR_PROMPT})

        response = client.messages.create(
            model="claude-sonnet-4-6",
            max_tokens=8192,
            messages=[{"role": "user", "content": content}],
        )
        all_lines.append(response.content[0].text.strip())

    return "\n".join(all_lines)


# ---------------------------------------------------------------------------
# PDF extraction
# ---------------------------------------------------------------------------

def _extract_pdf(pdf_path: str) -> str:
    try:
        import pdfplumber
    except ImportError:
        raise ImportError("pdfplumber not installed. Run: pip3 install pdfplumber")

    lines = []
    font_sizes = []
    meaningful_chars = 0

    # First pass: collect font sizes and count meaningful characters.
    # Dot-leader-only PDFs (TOC skeletons) have font_sizes but no real content.
    with pdfplumber.open(pdf_path) as pdf:
        for page in pdf.pages:
            words = page.extract_words(extra_attrs=["size", "fontname"])
            for w in words:
                s = w.get("size", 0)
                if s and s > 0:
                    font_sizes.append(s)
                text = w.get("text", "")
                meaningful_chars += sum(1 for c in text if c not in ".| \t\n")

    if not font_sizes or meaningful_chars < 80:
        logger.warning(
            "PDF appears image-based (%d meaningful chars). Falling back to Claude Vision OCR.",
            meaningful_chars,
        )
        return _extract_pdf_ocr(pdf_path)

    median_size = statistics.median(font_sizes)
    max_size = max(font_sizes)

    # Threshold: anything > median * 1.3 is a heading candidate
    # The largest font size is H1, next tier is H2, etc.
    heading_threshold = median_size * 1.3

    # Build size → heading level mapping from distinct large sizes
    large_sizes = sorted(set(s for s in font_sizes if s > heading_threshold), reverse=True)
    size_to_level = {}
    for i, s in enumerate(large_sizes[:4]):  # at most H1-H4
        size_to_level[s] = i + 1

    def _size_to_level(size: float) -> Optional[int]:
        if size <= heading_threshold:
            return None
        # Find closest size in map
        closest = min(size_to_level.keys(), key=lambda k: abs(k - size))
        if abs(closest - size) < 1.0:
            return size_to_level[closest]
        # Interpolate: bigger = higher level (lower number)
        for s, lvl in sorted(size_to_level.items(), reverse=True):
            if size >= s - 0.5:
                return lvl
        return None

    def _is_bold(fontname: str) -> bool:
        name = (fontname or "").lower()
        return "bold" in name or "-bd" in name or "heavy" in name or "black" in name

    with pdfplumber.open(pdf_path) as pdf:
        for page_num, page in enumerate(pdf.pages, 1):
            # Extract tables first, record their bboxes to skip those words
            tables = page.extract_tables()
            table_bboxes = [tbl_obj.bbox for tbl_obj in page.find_tables()] if hasattr(page, 'find_tables') else []

            # Extract words (characters grouped by proximity)
            words = page.extract_words(
                extra_attrs=["size", "fontname"],
                x_tolerance=3,
                y_tolerance=3,
            )

            if not words and not tables:
                continue

            # Group words into lines by y-position
            line_groups: dict = {}
            for w in words:
                # Skip words inside table bboxes
                wx = (w["x0"] + w["x1"]) / 2
                wy = (w["top"] + w["bottom"]) / 2
                in_table = any(
                    bx[0] <= wx <= bx[2] and bx[1] <= wy <= bx[3]
                    for bx in table_bboxes
                )
                if in_table:
                    continue
                y_key = round(w["top"] / 2) * 2  # bucket by 2pt
                if y_key not in line_groups:
                    line_groups[y_key] = []
                line_groups[y_key].append(w)

            # Build line objects sorted by vertical position
            line_objects = []
            for y_key in sorted(line_groups.keys()):
                words_in_line = sorted(line_groups[y_key], key=lambda w: w["x0"])
                text = " ".join(w["text"] for w in words_in_line).strip()
                if not text:
                    continue
                # Use the largest size in this line
                sizes = [w.get("size", 0) for w in words_in_line if w.get("size", 0) > 0]
                line_size = max(sizes) if sizes else 0
                # Bold: majority of words are bold
                bold_count = sum(1 for w in words_in_line if _is_bold(w.get("fontname", "")))
                line_bold = bold_count > len(words_in_line) / 2
                x0 = min(w["x0"] for w in words_in_line)
                line_objects.append({
                    "text": text, "size": line_size, "bold": line_bold,
                    "y": y_key, "x0": x0,
                })

            # Detect list items: indented lines starting with bullet or number
            page_margin = statistics.median([l["x0"] for l in line_objects]) if line_objects else 0
            _bullet_re = re.compile(r"^[•‣◦⁃∙\-\*\•]\s+|^\d+[\.\)]\s+")

            # Emit tagged lines + tables interleaved by vertical position
            # Build table entries with approximate y position
            table_entries = []
            if hasattr(page, 'find_tables'):
                for tbl_obj, tbl_data in zip(page.find_tables(), tables):
                    if not tbl_data:
                        continue
                    bbox = tbl_obj.bbox
                    y_pos = bbox[1]
                    table_entries.append({"y": y_pos, "data": tbl_data, "bbox": bbox})

            # Merge and sort by y
            all_items = []
            for lo in line_objects:
                all_items.append({"kind": "line", "y": lo["y"], "obj": lo})
            for te in table_entries:
                all_items.append({"kind": "table", "y": te["y"], "obj": te})
            all_items.sort(key=lambda x: x["y"])

            for item in all_items:
                if item["kind"] == "table":
                    tbl = item["obj"]["data"]
                    # Strip None cells, convert to strings
                    rows = []
                    for row in tbl:
                        clean_row = [str(c).strip() if c else "" for c in row]
                        rows.append(clean_row)
                    if not rows:
                        continue
                    nrows = len(rows)
                    ncols = max(len(r) for r in rows)
                    lines.append(f"[TABLE_START {nrows}x{ncols}]")
                    for row in rows:
                        lines.append("  | " + " | ".join(row) + " |")
                    lines.append("[TABLE_END]")
                else:
                    lo = item["obj"]
                    text = lo["text"]
                    size = lo["size"]
                    bold = lo["bold"]
                    x0 = lo["x0"]

                    level = _size_to_level(size)
                    if level is not None:
                        lines.append(f"[H{level}] {text}")
                    elif _bullet_re.match(text) or (x0 > page_margin + 15):
                        # Indented or bullet-prefixed → list item
                        clean = _bullet_re.sub("", text).strip()
                        lines.append(f"[LIST] {clean}")
                    elif bold:
                        lines.append(f"[BOLD] {text}")
                    else:
                        lines.append(f"[P] {text}")

    return "\n".join(lines)


# ---------------------------------------------------------------------------
# DOCX extraction (adapted from blog agent)
# ---------------------------------------------------------------------------

def _extract_docx(docx_path: str) -> str:
    try:
        from docx import Document
        from docx.oxml.ns import qn
        from docx.text.paragraph import Paragraph
        from docx.table import Table
    except ImportError:
        raise ImportError("python-docx not installed. Run: pip3 install python-docx")

    doc = Document(docx_path)
    lines = []

    for child in doc.element.body:
        tag = child.tag.split("}")[-1] if "}" in child.tag else child.tag

        if tag == "p":
            try:
                para = Paragraph(child, doc)
            except Exception:
                continue

            text = para.text.strip()

            if child.findall(".//" + qn("a:blip")):
                lines.append("[IMAGE_EMBEDDED]")
                if not text:
                    continue
            elif not text:
                continue

            urls = []
            for hl in para._element.findall(".//" + qn("w:hyperlink")):
                rel_id = hl.get(qn("r:id"))
                if rel_id:
                    try:
                        url = para.part.rels[rel_id].target_ref
                        if url:
                            urls.append(url)
                    except Exception:
                        pass

            style = para.style.name if para.style else "Normal"
            level_m = re.search(r"(\d+)", style) if style.startswith("Heading") else None

            if level_m:
                label = f"[H{level_m.group(1)}]"
            else:
                all_bold = all(r.bold for r in para.runs if r.text.strip())
                is_bold = all_bold and any(r.text.strip() for r in para.runs)
                num_pr = child.find(".//" + qn("w:numPr"))
                is_list = num_pr is not None or style in ("List Paragraph", "List Bullet", "List Number")
                if is_list:
                    label = "[LIST]"
                elif is_bold:
                    label = "[BOLD]"
                else:
                    label = "[P]"

            line = f"{label} {text}"
            if urls:
                line += f" [URLS: {', '.join(urls)}]"
            lines.append(line)

        elif tag == "tbl":
            try:
                tbl = Table(child, doc)
            except Exception:
                continue

            nrows = len(tbl.rows)
            ncols = len(tbl.rows[0].cells) if tbl.rows else 0

            if nrows == 1 and ncols == 1:
                cell = tbl.rows[0].cells[0]
                box_urls = []
                for p in cell.paragraphs:
                    for hl in p._element.findall(".//" + qn("w:hyperlink")):
                        rel_id = hl.get(qn("r:id"))
                        if rel_id:
                            try:
                                url = p.part.rels[rel_id].target_ref
                                if url:
                                    box_urls.append(url)
                            except Exception:
                                pass
                lines.append("[BOX_START]")
                for p in cell.paragraphs:
                    pt = p.text.strip()
                    if not pt:
                        continue
                    all_bold = all(r.bold for r in p.runs if r.text.strip())
                    pfx = "[BOLD]" if (all_bold and p.runs) else "[P]"
                    lines.append(f"  {pfx} {pt}")
                if box_urls:
                    lines.append(f"  [URLS: {', '.join(box_urls)}]")
                lines.append("[BOX_END]")
            else:
                lines.append(f"[TABLE_START {nrows}x{ncols}]")
                for row in tbl.rows:
                    cells = []
                    for c in row.cells:
                        parts = []
                        for p in c.paragraphs:
                            if p.text.strip():
                                parts.append(p.text.strip())
                            if p._element.findall(".//" + qn("a:blip")):
                                parts.append("[IMAGE_EMBEDDED]")
                        if not parts and c._element.findall(".//" + qn("a:blip")):
                            parts.append("[IMAGE_EMBEDDED]")
                        cells.append(" ".join(parts) if parts else "")
                    lines.append("  | " + " | ".join(cells) + " |")
                lines.append("[TABLE_END]")

    return "\n".join(lines)


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------

def extract(file_path: str) -> str:
    """Route PDF or DOCX to the appropriate extractor. Returns tagged plaintext."""
    path = Path(file_path)
    ext = path.suffix.lower()
    if ext == ".pdf":
        logger.info(f"PDF mode: {path.name}")
        return _extract_pdf(file_path)
    elif ext in (".docx", ".doc"):
        logger.info(f"DOCX mode: {path.name}")
        return _extract_docx(file_path)
    else:
        raise ValueError(f"Unsupported file type: {ext} — expected .pdf or .docx")


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO, format="%(levelname)s: %(message)s")
    parser = argparse.ArgumentParser(description="Parse a PDF or DOCX whitepaper to tagged plaintext.")
    parser.add_argument("file_path", help="Path to .pdf or .docx file")
    parser.add_argument("--output", "-o", default=None, help="Output path for tagged text (default: print to stdout)")
    args = parser.parse_args()

    text = extract(args.file_path)

    if args.output:
        Path(args.output).parent.mkdir(parents=True, exist_ok=True)
        Path(args.output).write_text(text, encoding="utf-8")
        print(f"Written: {args.output}")
        print(f"Lines: {len(text.splitlines())}")
    else:
        print(text)
